"""Builds the FR -> EN / BM dictionary of the muso port from the CHT sources (no translation
of our own): converted forms (out/*.i18n.json), contact forms, messages-*.properties and
the form titles of *.properties.json. Writes dictionary.json: { fr: { en, bm } }.

Usage: python3 dictionary.py <muso-mali>
"""
import glob
import json
import os
import re
import sys

import openpyxl

SOURCE = sys.argv[1]
HERE = os.path.dirname(os.path.abspath(__file__))
pairs = {}


def clean(text):
    text = re.sub(r'<i[^>]*>\s*</i>', '', str(text), flags=re.I)
    text = re.sub(r'<[^>]+>', '', text)
    text = re.sub(r'[ \t]+', ' ', text).strip()
    return text


def template(text):
    # Same transformation as convert.mjs for ${name} in labels.
    return re.sub(r'\$\{([^}]+)\}', lambda m: '{app.currentForm[' + json.dumps(m.group(1).strip()) + "] ?? ''}", text)


def add(fr, en, bm):
    if not fr or str(fr).strip() in ('', 'NO_LABEL'):
        return
    for key in {clean(fr), template(clean(fr))}:
        entry = pairs.setdefault(key, {})
        if en and 'en' not in entry:
            entry['en'] = template(clean(en)) if '${' in str(en) else clean(en)
        if bm and 'bm' not in entry:
            entry['bm'] = template(clean(bm)) if '${' in str(bm) else clean(bm)


# 1. Converted app forms.
for path in glob.glob(os.path.join(HERE, 'out', '*.i18n.json')):
    data = json.load(open(path, encoding='utf-8'))
    for fr, en in data.get('en', {}).items():
        add(fr, en, data.get('bm', {}).get(fr))
    for fr, bm in data.get('bm', {}).items():
        add(fr, data.get('en', {}).get(fr), bm)

# 2. Contact forms (attributes of persons and places) and every app form again (choices).
for path in glob.glob(os.path.join(SOURCE, 'forms', '*', '*.xlsx')):
    workbook = openpyxl.load_workbook(path, read_only=True)
    for sheet_name in ('survey', 'choices'):
        if sheet_name not in workbook.sheetnames:
            continue
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else '' for h in rows[0]]
        for base in ('label', 'hint', 'constraint_message', 'required_message'):
            cols = {lang: header.index(f'{base}::{lang}') for lang in ('fr', 'en', 'bm') if f'{base}::{lang}' in header}
            if 'fr' not in cols:
                continue
            for row in rows[1:]:
                get = lambda lang: row[cols[lang]] if lang in cols and cols[lang] < len(row) else None
                add(get('fr'), get('en'), get('bm'))

# 3. messages-*.properties, matched by key.
def properties(lang):
    result = {}
    path = os.path.join(SOURCE, 'translations', f'messages-{lang}.properties')
    for line in open(path, encoding='utf-8'):
        if '=' in line and not line.lstrip().startswith('#'):
            key, value = line.split('=', 1)
            result[key.strip()] = value.strip()
    return result

fr_msgs, en_msgs, bm_msgs = properties('fr'), properties('en'), properties('bm')
for key, fr in fr_msgs.items():
    add(fr, en_msgs.get(key), bm_msgs.get(key))

# 4. Form titles.
for path in glob.glob(os.path.join(SOURCE, 'forms', '*', '*.properties.json')):
    titles = {t['locale']: t['content'] for t in json.load(open(path, encoding='utf-8')).get('title', [])}
    add(titles.get('fr'), titles.get('en'), titles.get('bm'))

json.dump(pairs, open(os.path.join(HERE, 'dictionary.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(len(pairs), 'FR texts,', sum('en' in v for v in pairs.values()), 'with EN,', sum('bm' in v for v in pairs.values()), 'with BM')
