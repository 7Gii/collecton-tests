"""Exports the CHT XLSForms of muso-mali to JSON (survey, choices, settings) for convert.mjs.

Usage: python3 extract.py <muso-mali/forms/app> <form> [<form> ...]
"""
import json
import os
import sys

import openpyxl


def sheet(workbook, name):
    if name not in workbook.sheetnames:
        return []
    rows = list(workbook[name].iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else None for h in rows[0]]
    result = []
    for row in rows[1:]:
        if not row or all(cell is None for cell in row):
            continue
        record = {}
        for i in range(min(len(header), len(row))):
            value = row[i]
            if not header[i] or value is None or str(value).strip() == '':
                continue
            record[header[i]] = int(value) if isinstance(value, float) and value.is_integer() else value
        result.append(record)
    return result


def main():
    source, forms = sys.argv[1], sys.argv[2:]
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'json')
    os.makedirs(out, exist_ok=True)
    for form in forms:
        workbook = openpyxl.load_workbook(os.path.join(source, f'{form}.xlsx'), read_only=True)
        data = {name: sheet(workbook, name) for name in ('survey', 'choices', 'settings')}
        with open(os.path.join(out, f'{form}.json'), 'w', encoding='utf-8') as handle:
            json.dump(data, handle, ensure_ascii=False, indent=1, default=str)
        print(form, len(data['survey']), 'rows')


if __name__ == '__main__':
    main()
